import pandas as pd
from openpyxl import load_workbook

agg = pd.read_excel("module_b\Агрегация по районам.xlsx")

#XLSX, выгруженный из ритм

file_name = "module_b\Районы.xlsx"

wb = load_workbook(file_name)
ws = wb["Районы"]

# Нудно найти системные названия значений: Популяции и Рабочих мест (Найти их айди в таблице)
# Дальше циклом обновить

# BL - Место с workplaces
# BO - Место с population

for _, row in agg.iterrows():
	zone_id = int(row["no"])

	for excel_row in range(3, ws.max_row + 1):
		#Проверяем первый столбик
		if ws.cell(excel_row, 1).value == zone_id:
			ws[f"BL{excel_row}"] = row["workplaces"]
			ws[f"BO{excel_row}"] = row["population"]
			break

wb.save("Районы для импорта.xlsx")
